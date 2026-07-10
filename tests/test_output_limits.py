import unittest
from collections import Counter
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import load_workbook

from splitter import OutputManager


class OutputWorksheetLimitTests(unittest.TestCase):
    def test_priority_column_cannot_cross_excel_column_limit(self):
        with TemporaryDirectory() as tmp, patch(
            'splitter.EXCEL_MAX_COLUMNS', 4
        ):
            with self.assertRaisesRegex(
                ValueError,
                r"4 input columns.*Priority.*4-column limit",
            ):
                OutputManager(
                    Path(tmp),
                    'crawl',
                    ['Type', 'Source', 'Destination', 'Source Segments'],
                    Counter(),
                    no_summary=True,
                )

    def test_data_rolls_over_before_worksheet_row_limit(self):
        with TemporaryDirectory() as tmp, patch(
            'splitter.EXCEL_MAX_ROWS', 4, create=True
        ):
            output_dir = Path(tmp)
            manager = OutputManager(
                output_dir,
                'crawl',
                ['Source Segments', 'Destination'],
                Counter(),
                no_summary=True,
            )
            for row_number in range(7):
                manager.append(
                    'USA',
                    ('USA', f'https://example.com/{row_number}'),
                    f'https://example.com/{row_number}',
                    'LOW',
                )
            manager.close_all()

            workbook = load_workbook(output_dir / 'crawl_USA.xlsx', read_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ['Data', 'Data 2', 'Data 3'])
                row_counts = [
                    len(list(workbook[sheet].iter_rows(values_only=True)))
                    for sheet in workbook.sheetnames
                ]
                self.assertEqual(row_counts, [4, 4, 2])
                data_rows = sum(count - 1 for count in row_counts)
                self.assertEqual(data_rows, 7)
                for sheet in workbook.worksheets:
                    header = next(sheet.iter_rows(values_only=True))
                    self.assertEqual(
                        header,
                        ('Source Segments', 'Destination', 'Priority'),
                    )
                written_rows = [
                    row
                    for sheet in workbook.worksheets
                    for row in list(sheet.iter_rows(values_only=True))[1:]
                ]
                expected_rows = [
                    (
                        'USA',
                        f'https://example.com/{row_number}',
                        'LOW',
                    )
                    for row_number in range(7)
                ]
                self.assertEqual(written_rows, expected_rows)
            finally:
                workbook.close()

    def test_summary_rolls_over_before_worksheet_row_limit(self):
        destinations = [f'https://example.com/{number}' for number in range(7)]
        with TemporaryDirectory() as tmp, patch(
            'splitter.EXCEL_MAX_ROWS', 4, create=True
        ):
            output_dir = Path(tmp)
            manager = OutputManager(
                output_dir,
                'crawl',
                ['Source Segments', 'Destination'],
                Counter({destination: 1 for destination in destinations}),
                no_summary=False,
            )
            for destination in destinations:
                manager.append(
                    'USA',
                    ('USA', destination),
                    destination,
                    'LOW',
                )
            manager.close_all()

            workbook = load_workbook(output_dir / 'crawl_USA.xlsx', read_only=True)
            try:
                summary_names = [
                    name for name in workbook.sheetnames if name.startswith('Summary')
                ]
                self.assertEqual(summary_names, ['Summary', 'Summary 2', 'Summary 3'])
                row_counts = [
                    len(list(workbook[name].iter_rows(values_only=True)))
                    for name in summary_names
                ]
                self.assertEqual(row_counts, [4, 4, 2])
                written_destinations = {
                    row[1]
                    for name in summary_names
                    for row in list(workbook[name].iter_rows(values_only=True))[1:]
                }
                self.assertEqual(written_destinations, set(destinations))
            finally:
                workbook.close()


if __name__ == '__main__':
    unittest.main()
