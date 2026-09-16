import json
import unittest
from argparse import Namespace
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

import splitter
from splitter import _resolve_buckets, analyze_workbook, split_workbook


class RegionalSubsetConfigTests(unittest.TestCase):
    def test_directional_config_validation(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / 'config.json'
            path.write_text(json.dumps({'To / destination': {'EU': ['/folder/']}}))
            self.assertEqual(splitter.load_regional_subset_config(path), {
                'From / source': {}, 'To / destination': {'EU': ('folder',)},
            })
            for config in (
                {'To / destination': []},
                {'To / destination': {'EU': ['../folder']}},
                {'To / destination': {'EU': ['folder', 'Folder']}},
                {'To / destination': {}, 'EU': []},
                {'From / source': {'EU': ['to_folder']},
                 'To / destination': {'EU': ['folder']}},
            ):
                with self.subTest(config=config):
                    path.write_text(json.dumps(config))
                    with self.assertRaises(ValueError):
                        splitter.load_regional_subset_config(path)

    def test_missing_config_preserves_existing_behavior(self):
        with TemporaryDirectory() as tmp:
            config_path = Path(tmp) / 'missing.json'

            self.assertEqual(
                splitter.load_regional_subset_config(config_path),
                {},
            )

    def test_loads_and_normalizes_region_path_lists(self):
        with TemporaryDirectory() as tmp:
            config_path = Path(tmp) / 'subsets.json'
            config_path.write_text(
                json.dumps({
                    'EU': ['/campaign/', 'shipping/surcharges'],
                    'USA': [],
                }),
                encoding='utf-8',
            )

            loader = getattr(splitter, 'load_regional_subset_config', None)
            self.assertIsNotNone(loader)
            subsets = loader(config_path)

            self.assertEqual(
                subsets,
                {'EU': ('campaign', 'shipping/surcharges')},
            )

    def test_rejects_unknown_regions(self):
        with TemporaryDirectory() as tmp:
            config_path = Path(tmp) / 'subsets.json'
            config_path.write_text(
                json.dumps({'EMEA': ['campaign']}),
                encoding='utf-8',
            )

            with self.assertRaisesRegex(ValueError, 'unknown region.*EMEA'):
                splitter.load_regional_subset_config(config_path)

    def test_rejects_duplicate_paths_case_insensitively(self):
        with TemporaryDirectory() as tmp:
            config_path = Path(tmp) / 'subsets.json'
            config_path.write_text(
                json.dumps({'EU': ['campaign', '/Campaign/']}),
                encoding='utf-8',
            )

            with self.assertRaisesRegex(ValueError, 'duplicate path.*Campaign'):
                splitter.load_regional_subset_config(config_path)

    def test_rejects_non_object_configuration(self):
        with TemporaryDirectory() as tmp:
            config_path = Path(tmp) / 'subsets.json'
            config_path.write_text('[]', encoding='utf-8')

            with self.assertRaisesRegex(ValueError, 'JSON object'):
                splitter.load_regional_subset_config(config_path)

    def test_rejects_non_list_region_paths(self):
        with TemporaryDirectory() as tmp:
            config_path = Path(tmp) / 'subsets.json'
            config_path.write_text(
                json.dumps({'EU': 'campaign'}),
                encoding='utf-8',
            )

            with self.assertRaisesRegex(ValueError, 'EU.*array'):
                splitter.load_regional_subset_config(config_path)

    def test_rejects_unsafe_url_path_entries(self):
        invalid_paths = (
            'https://www.fedex.com/campaign',
            'campaign?preview=true',
            '../campaign',
        )
        for invalid_path in invalid_paths:
            with self.subTest(path=invalid_path), TemporaryDirectory() as tmp:
                config_path = Path(tmp) / 'subsets.json'
                config_path.write_text(
                    json.dumps({'EU': [invalid_path]}),
                    encoding='utf-8',
                )

                with self.assertRaisesRegex(ValueError, 'unsafe path'):
                    splitter.load_regional_subset_config(config_path)

    def test_rejects_paths_that_would_share_an_output_filename(self):
        with TemporaryDirectory() as tmp:
            config_path = Path(tmp) / 'subsets.json'
            config_path.write_text(
                json.dumps({'EU': ['shipping/surcharges', 'shipping_surcharges']}),
                encoding='utf-8',
            )

            with self.assertRaisesRegex(ValueError, 'same output name'):
                splitter.load_regional_subset_config(config_path)


class RegionalSubsetRoutingTests(unittest.TestCase):
    def test_destination_subsets_across_reordered_sheets(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / 'crawl.xlsx'
            workbook = Workbook()
            headers = ['Type', 'Source', 'Destination', 'Source Segments']
            destination = 'https://www.fedex.com/nl-nl/FOLDER/tracking.html'
            rows = [
                ['Hyperlink', 'https://www.fedex.com/nl-nl/customer-support', destination, 'EU'],
                ['Hyperlink', 'https://www.fedex.com/nl-nl/folder', 'https://example.com/elsewhere', 'EU'],
                ['Hyperlink', 'https://example.com/page', 'https://example.com/folderish/page?x=/folder/', 'EU'],
                ['XML Sitemap', 'https://example.com/page', destination, 'EU'],
            ]
            sheet = workbook.active
            sheet.append(headers)
            for row in rows:
                sheet.append(row)
            sheet = workbook.create_sheet('Reordered')
            sheet.append(list(reversed(headers)))
            sheet.append(list(reversed(rows[0])))
            workbook.save(path)
            config_path = root / 'config.json'
            config_path.write_text(json.dumps({
                'From / source': {'EU': ['folder']},
                'To / destination': {'EU': ['folder', 'folder/tracking.html']},
            }))
            args = Namespace(split='region', url_depth=2, url_pattern=None,
                             no_filter=False, no_summary=False,
                             regional_subsets=splitter.load_regional_subset_config(config_path))
            canonical, counts, buckets, _, _ = analyze_workbook(path, args)
            self.assertEqual(dict(buckets), {
                'EU': 4, 'EU_folder': 1, 'EU_to_folder': 2,
                'EU_to_folder_tracking.html': 2,
            })
            output = root / 'output'
            output.mkdir()
            self.assertEqual(split_workbook(path, args, canonical, counts, output), 4)
            result = load_workbook(output / 'crawl_EU_to_folder.xlsx', read_only=True)
            try:
                data = list(result['Data'].values)
                self.assertEqual(len(data), 3)
                self.assertTrue(all(row[2] == destination for row in data[1:]))
                summary = list(result['Summary'].values)
                self.assertEqual(summary[1][1:3], (destination, 2))
            finally:
                result.close()

    def test_active_subset_config_requires_source_url_column(self):
        with TemporaryDirectory() as tmp:
            xlsx_path = Path(tmp) / 'crawl.xlsx'
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(['Type', 'Destination', 'Source Segments'])
            sheet.append(['Hyperlink', 'https://example.com/missing', 'EU'])
            workbook.save(xlsx_path)

            args = Namespace(
                split='region',
                url_depth=2,
                url_pattern=None,
                no_filter=False,
                no_summary=False,
                regional_subsets={'EU': ('campaign',)},
            )

            with self.assertRaisesRegex(ValueError, 'Source'):
                analyze_workbook(xlsx_path, args)

    def test_region_rows_keep_base_bucket_and_gain_every_matching_subset(self):
        row = (
            'https://www.fedex.com/en-gb/campaign/shipping/surcharges/index.html',
            'EU',
        )

        try:
            buckets = _resolve_buckets(
                row,
                seg_col=1,
                source_col=0,
                split_mode='region',
                url_depth=2,
                url_pattern=None,
                regional_subsets={
                    'EU': ('campaign', 'shipping/surcharges'),
                },
            )
        except TypeError as error:
            self.fail(f'regional subset routing is unavailable: {error}')

        self.assertEqual(
            buckets,
            ['EU', 'EU_campaign', 'EU_shipping_surcharges'],
        )

    def test_configured_subsets_do_not_change_url_split_mode(self):
        row = ('https://www.fedex.com/en-gb/campaign/page.html', 'EU')

        buckets = _resolve_buckets(
            row,
            seg_col=1,
            source_col=0,
            split_mode='url',
            url_depth=1,
            url_pattern=None,
            regional_subsets={'EU': ('campaign',)},
        )

        self.assertEqual(buckets, ['campaign'])

    def test_workbook_writes_base_region_and_matching_subset_files(self):
        with TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            xlsx_path = tmp_path / 'crawl.xlsx'
            output_dir = tmp_path / 'output'
            output_dir.mkdir()

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = 'All Inlinks'
            sheet.append(['Type', 'Source', 'Destination', 'Source Segments'])
            sheet.append([
                'Hyperlink',
                'https://www.fedex.com/en-gb/campaign/sale.html',
                'https://example.com/destination-1',
                'EU',
            ])
            sheet.append([
                'Hyperlink',
                'https://www.fedex.com/fr-fr/shipping/surcharges/fuel.html',
                'https://example.com/destination-2',
                'EU',
            ])
            sheet.append([
                'Hyperlink',
                'https://www.fedex.com/de-de/services.html',
                'https://example.com/destination-3',
                'EU',
            ])
            workbook.save(xlsx_path)

            args = Namespace(
                split='region',
                url_depth=2,
                url_pattern=None,
                no_filter=False,
                no_summary=False,
                regional_subsets={
                    'EU': ('campaign', 'shipping/surcharges'),
                },
            )

            headers, dest_counter, bucket_counter, _, _ = analyze_workbook(
                xlsx_path, args
            )
            files_created = split_workbook(
                xlsx_path, args, headers, dest_counter, output_dir
            )

            self.assertEqual(
                dict(bucket_counter),
                {'EU': 3, 'EU_campaign': 1, 'EU_shipping_surcharges': 1},
            )
            self.assertEqual(files_created, 3)

            expected_rows = {
                'crawl_EU.xlsx': 3,
                'crawl_EU_campaign.xlsx': 1,
                'crawl_EU_shipping_surcharges.xlsx': 1,
            }
            for filename, expected_count in expected_rows.items():
                result = load_workbook(output_dir / filename, read_only=True)
                try:
                    rows = list(result['Data'].iter_rows(values_only=True))
                    self.assertEqual(len(rows) - 1, expected_count)
                finally:
                    result.close()


if __name__ == '__main__':
    unittest.main()
